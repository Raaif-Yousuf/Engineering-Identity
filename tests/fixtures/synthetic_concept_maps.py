"""Synthetic concept-map workbook builder used only by tests.

Never touches or embeds real student data: every workbook this module
writes uses letter-coded, made-up concept labels and a fabricated
participant code.
"""

from pathlib import Path

import openpyxl


def write_concept_map_workbook(
    path: Path,
    participant_code: str,
    edges: list[tuple[str, str]],
    include_name_sheet: bool = True,
    source_sheet_name: str = "Sources",
    sink_sheet_name: str = "Sinks",
    trailing_blank_rows: int = 0,
) -> Path:
    """Write a minimal synthetic concept-map workbook.

    Mirrors the real workbook format: a `source_sheet_name` sheet and a
    `sink_sheet_name` sheet, each a single column, row i giving one
    directed relation `edges[i][0] -> edges[i][1]`. `trailing_blank_rows`
    appends empty padding rows after the real data, matching how real
    workbooks pad both sheets to a fixed template length.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    src_ws = wb.create_sheet(source_sheet_name)
    snk_ws = wb.create_sheet(sink_sheet_name)
    for source, sink in edges:
        src_ws.append([source])
        snk_ws.append([sink])
    for _ in range(trailing_blank_rows):
        src_ws.append([None])
        snk_ws.append([None])

    if include_name_sheet:
        name_ws = wb.create_sheet("Name")
        name_ws["A1"] = participant_code

    wb.save(path)
    return path
